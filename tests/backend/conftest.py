import os
import tempfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import delete

# Tests use a dedicated temporary database, never the personal master.
os.environ["STUDIO_DATA_DIR"] = tempfile.mkdtemp(prefix="reshi-studio-tests-")
os.environ["AI_MODE"] = "disabled"
os.environ.pop("RESUME_DOCX", None)
os.environ.pop("RESUME_PDF", None)

from backend.database import Export, Job, Meta, Revision, SessionLocal, Version
from backend.main import app


def make_docx(path: Path):
    document = Document()
    document.add_paragraph("Test Candidate", "Title")
    document.add_paragraph("QA Engineer | SDET", "Subtitle")
    document.add_paragraph("Bengaluru, India | +91 9999999999 | candidate@example.test")
    document.add_paragraph("example.com")
    document.add_heading("PROFESSIONAL SUMMARY", 1)
    document.add_paragraph(
        "Built Selenium test frameworks. Developed Python APIs. Delivered 100+ regression tests."
    )
    document.add_heading("SKILLS", 1)
    document.add_paragraph("Languages: Java, Python, JavaScript")
    document.add_paragraph("Test automation: Selenium, TestNG, Robot Framework")
    document.add_heading("WORK EXPERIENCE", 1)
    document.add_paragraph("Example Company\tJan 2023 - Present")
    document.add_paragraph("QA Engineer | Bengaluru")
    document.add_paragraph("Built Selenium automation and 100+ regression tests.", "List Bullet")
    document.add_paragraph("Developed Python APIs with deterministic validation.", "List Bullet")
    document.add_heading("PROJECTS", 1)
    document.add_heading("Test Engine | Automated validation", 2)
    document.add_paragraph("Python, pytest")
    document.add_paragraph("Built a Python testing engine.", "List Bullet")
    document.add_heading("AWARDS AND HONORS", 1)
    document.add_paragraph("Test Hackathon Winner")
    document.add_heading("CERTIFICATIONS", 1)
    document.add_paragraph("Test Credential")
    document.add_heading("EDUCATION", 1)
    document.add_paragraph("Bachelor of Engineering")
    document.add_paragraph("Example College")
    document.add_paragraph("2019 - 2022")
    document.save(path)


@pytest.fixture(scope="session")
def client():
    with TestClient(app) as client:
        yield client


@pytest.fixture(autouse=True)
def reset_database(client):
    with SessionLocal.begin() as db:
        for table in (Export, Version, Job, Meta, Revision):
            db.execute(delete(table))


@pytest.fixture
def docx_path(tmp_path):
    path = tmp_path / "canonical.docx"
    make_docx(path)
    return path


@pytest.fixture
def imported(client, docx_path):
    with docx_path.open("rb") as stream:
        result = client.post("/api/import", files={"docx": ("canonical.docx", stream)})
    assert result.status_code == 200, result.text
    return result.json()
